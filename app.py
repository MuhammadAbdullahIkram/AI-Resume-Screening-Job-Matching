
import re
from io import BytesIO

import pandas as pd
import streamlit as st
from pypdf import PdfReader
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


# ==================================================
# PAGE CONFIGURATION
# ==================================================

st.set_page_config(
    page_title="AI Resume Matcher",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded"
)


# ==================================================
# BLUE UI THEME
# ==================================================

st.markdown("""
<style>
    .stApp {
        background-color: #F4F7FC;
    }

    [data-testid="stHeader"] {
        background-color: #F4F7FC;
    }

    .main .block-container {
        max-width: 1200px;
        padding-top: 2rem;
        padding-bottom: 3rem;
    }

    .hero {
        background: linear-gradient(120deg, #123B78, #2563EB);
        padding: 30px;
        border-radius: 18px;
        color: white;
        margin-bottom: 25px;
    }

    .hero h1 {
        color: white;
        font-size: 34px;
        margin-bottom: 8px;
    }

    .hero p {
        color: #DBEAFE;
        font-size: 16px;
        margin-bottom: 0;
    }

    h2, h3 {
        color: #173B70;
    }

    div.stButton > button[kind="primary"] {
        background-color: #2563EB;
        color: white;
        border: none;
        border-radius: 9px;
        min-height: 45px;
        font-weight: 600;
    }

    div.stButton > button[kind="primary"]:hover {
        background-color: #1D4ED8;
        color: white;
    }

    div[data-testid="stMetric"] {
        background-color: white;
        border: 1px solid #DCE6F5;
        padding: 18px;
        border-radius: 12px;
    }

    section[data-testid="stSidebar"] {
        background-color: #EAF1FC;
    }

    .note {
        background-color: #E0ECFF;
        border-left: 4px solid #2563EB;
        padding: 14px;
        border-radius: 7px;
        color: #173B70;
        margin: 10px 0 20px 0;
    }
</style>
""", unsafe_allow_html=True)


# ==================================================
# HEADER
# ==================================================

st.markdown("""
<div class="hero">
    <h1>📄 AI Resume Matcher</h1>
    <p>Smart Resume Screening & Job Matching System</p>
</div>
""", unsafe_allow_html=True)

st.write(
    "Compare resumes with job requirements using text similarity "
    "and skill keyword matching."
)


# ==================================================
# SIDEBAR
# ==================================================

with st.sidebar:
    st.title("⚙️ Settings")
    st.markdown("---")

    st.subheader("Features")
    st.write("📄 CSV resume upload")
    st.write("📕 PDF resume upload")
    st.write("📝 Word document upload")
    st.write("🔎 Skill matching")
    st.write("📊 Candidate ranking")
    st.write("⬇️ CSV results export")

    st.markdown("---")
    st.caption(
        "Built with Python, Pandas, Scikit-learn and Streamlit."
    )


# ==================================================
# FILE TEXT EXTRACTION
# ==================================================

def extract_resume_files(files):
    all_resumes = []
    errors = []

    for uploaded_file in files:
        filename = uploaded_file.name
        extension = filename.lower().rsplit(".", 1)[-1]

        try:

            # CSV: multiple candidates
            if extension == "csv":
                uploaded_file.seek(0)

                try:
                    df_csv = pd.read_csv(
                        uploaded_file,
                        encoding="utf-8-sig"
                    )
                except UnicodeDecodeError:
                    uploaded_file.seek(0)
                    df_csv = pd.read_csv(
                        uploaded_file,
                        encoding="latin-1"
                    )

                df_csv.columns = (
                    df_csv.columns.astype(str)
                    .str.strip()
                    .str.lower()
                )

                required_columns = {
                    "candidate_id",
                    "resume_text"
                }

                if not required_columns.issubset(df_csv.columns):
                    errors.append(
                        f"{filename}: CSV must contain "
                        "'candidate_id' and 'resume_text' columns."
                    )
                    continue

                for _, row in df_csv.iterrows():
                    if pd.isna(row["resume_text"]):
                        continue

                    resume_text = str(
                        row["resume_text"]
                    ).strip()

                    if not resume_text:
                        continue

                    candidate_id = (
                        str(row["candidate_id"]).strip()
                        if pd.notna(row["candidate_id"])
                        else "Unknown"
                    )

                    all_resumes.append({
                        "candidate_id": candidate_id,
                        "resume_text": resume_text,
                        "source_file": filename
                    })

            # PDF: one candidate per PDF
            elif extension == "pdf":
                pdf_reader = PdfReader(
                    BytesIO(uploaded_file.getvalue())
                )

                if pdf_reader.is_encrypted:
                    try:
                        result = pdf_reader.decrypt("")
                        if not result:
                            errors.append(
                                f"{filename}: Password-protected PDF."
                            )
                            continue
                    except Exception:
                        errors.append(
                            f"{filename}: Could not open encrypted PDF."
                        )
                        continue

                pages = []

                for page in pdf_reader.pages:
                    pages.append(page.extract_text() or "")

                resume_text = "\n".join(pages).strip()

                if not resume_text:
                    errors.append(
                        f"{filename}: No selectable text found. "
                        "This may be a scanned PDF requiring OCR."
                    )
                    continue

                all_resumes.append({
                    "candidate_id": filename.rsplit(".", 1)[0],
                    "resume_text": resume_text,
                    "source_file": filename
                })

            # Word: DOCX format
            elif extension == "docx":
                document = Document(
                    BytesIO(uploaded_file.getvalue())
                )

                paragraphs = [
                    paragraph.text.strip()
                    for paragraph in document.paragraphs
                    if paragraph.text.strip()
                ]

                # Also collect text from Word tables
                for table in document.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            cell_text = cell.text.strip()
                            if cell_text:
                                paragraphs.append(cell_text)

                resume_text = "\n".join(paragraphs).strip()

                if not resume_text:
                    errors.append(
                        f"{filename}: No readable text found."
                    )
                    continue

                all_resumes.append({
                    "candidate_id": filename.rsplit(".", 1)[0],
                    "resume_text": resume_text,
                    "source_file": filename
                })

        except Exception as e:
            errors.append(
                f"{filename}: Could not process this file. {str(e)}"
            )

    return pd.DataFrame(all_resumes), errors


# ==================================================
# JOB DESCRIPTION
# ==================================================

st.subheader("1. Job Requirements")

job_description = st.text_area(
    "Job Description",
    placeholder=(
        "Example: We need a Data Scientist with Python, SQL, "
        "Machine Learning, Pandas and Data Analysis skills."
    ),
    height=150
)

skills_input = st.text_input(
    "Required Skills (comma-separated)",
    value="Python, SQL, Machine Learning, Pandas, Data Analysis"
)


# ==================================================
# FILE UPLOAD
# ==================================================

st.subheader("2. Upload Candidate Resumes")

st.caption(
    "Supported formats: CSV, PDF and Word (.docx). "
    "You can select multiple files."
)

uploaded_files = st.file_uploader(
    "Choose resume files",
    type=["csv", "pdf", "docx"],
    accept_multiple_files=True,
    help=(
        "CSV files must contain candidate_id and resume_text columns. "
        "Each PDF or DOCX file is treated as one candidate."
    )
)

if uploaded_files:
    st.success(
        f"{len(uploaded_files)} file(s) selected."
    )

st.markdown("""
<div class="note">
<strong>Privacy & responsible use:</strong>
Use fictional or consented resume data. Matching scores are
approximate text and keyword indicators, not measures of a
person's overall ability or suitability. Review candidates fairly
and make decisions through human assessment.
</div>
""", unsafe_allow_html=True)


# ==================================================
# ANALYZE RESUMES
# ==================================================

if st.button(
    "🔍 Analyze Resumes",
    type="primary",
    use_container_width=True
):

    if not job_description.strip():
        st.error("Please enter a job description.")
        st.stop()

    if not uploaded_files:
        st.error("Please upload at least one resume file.")
        st.stop()

    with st.spinner("Reading files and analyzing resumes..."):

        df, file_errors = extract_resume_files(
            uploaded_files
        )

    for error in file_errors:
        st.warning(error)

    if df.empty:
        st.error(
            "No readable resumes were found. Check the file formats "
            "and make sure the documents contain selectable text."
        )
        st.stop()

    st.success(
        f"Successfully extracted {len(df)} resume(s)."
    )

    # ----------------------------------------------
    # PREPARE SKILLS
    # ----------------------------------------------

    skills = [
        skill.strip().lower()
        for skill in skills_input.split(",")
        if skill.strip()
    ]

    resume_texts = df["resume_text"].astype(str).tolist()

    # ----------------------------------------------
    # TF-IDF TEXT SIMILARITY
    # ----------------------------------------------

    try:
        vectorizer = TfidfVectorizer(
            stop_words="english"
        )

        matrix = vectorizer.fit_transform(
            [job_description] + resume_texts
        )

        similarities = cosine_similarity(
            matrix[0:1],
            matrix[1:]
        ).flatten()

    except ValueError:
        st.error(
            "The text could not be compared. Please provide "
            "more meaningful text in the job description and resumes."
        )
        st.stop()

    # ----------------------------------------------
    # SKILL COVERAGE
    # ----------------------------------------------

    def calculate_skill_coverage(text):
        text = text.lower()

        matched_skills = [
            skill for skill in skills
            if re.search(
                r"(?<!\w)" + re.escape(skill) + r"(?!\w)",
                text
            )
        ]

        if skills:
            coverage = (
                len(matched_skills) / len(skills)
            ) * 100
        else:
            coverage = 0.0

        return pd.Series({
            "skill_coverage_pct": coverage,
            "matched_skills": ", ".join(matched_skills)
        })

    skill_results = df["resume_text"].apply(
        calculate_skill_coverage
    )

    df["text_similarity_pct"] = similarities * 100

    df["skill_coverage_pct"] = (
        skill_results["skill_coverage_pct"].values
    )

    df["matched_skills"] = (
        skill_results["matched_skills"].values
    )

    # ----------------------------------------------
    # COMBINED SCORE
    # ----------------------------------------------

    df["combined_score"] = (
        0.70 * df["text_similarity_pct"]
        + 0.30 * df["skill_coverage_pct"]
    )

    df = df.sort_values(
        "combined_score",
        ascending=False
    ).reset_index(drop=True)

    df.insert(
        0,
        "rank",
        range(1, len(df) + 1)
    )

    # ----------------------------------------------
    # RESULTS SUMMARY
    # ----------------------------------------------

    st.markdown("---")
    st.subheader("3. Matching Results")

    col1, col2, col3 = st.columns(3)

    col1.metric(
        "Resumes Analyzed",
        len(df)
    )

    col2.metric(
        "Highest Match Score",
        f"{df.iloc[0]['combined_score']:.1f}%"
    )

    col3.metric(
        "Average Skill Coverage",
        f"{df['skill_coverage_pct'].mean():.1f}%"
    )

    # ----------------------------------------------
    # RANKING TABLE
    # ----------------------------------------------

    st.subheader("Candidate Ranking")

    display_df = df[
        [
            "rank",
            "candidate_id",
            "source_file",
            "text_similarity_pct",
            "skill_coverage_pct",
            "combined_score",
            "matched_skills"
        ]
    ].copy()

    st.dataframe(
        display_df.style.format({
            "text_similarity_pct": "{:.1f}%",
            "skill_coverage_pct": "{:.1f}%",
            "combined_score": "{:.1f}%"
        }),
        use_container_width=True,
        hide_index=True
    )

    # ----------------------------------------------
    # DOWNLOAD RESULTS
    # ----------------------------------------------

    csv_data = df.to_csv(
        index=False
    ).encode("utf-8-sig")

    st.download_button(
        label="⬇️ Download Ranking CSV",
        data=csv_data,
        file_name="resume_ranking.csv",
        mime="text/csv",
        use_container_width=True
    )

    # ----------------------------------------------
    # RESPONSIBLE USE
    # ----------------------------------------------

    st.caption(
        "Combined score = 70% text similarity + 30% skill coverage. "
        "This is a basic keyword-based prototype, not a validated "
        "hiring assessment. Scores should not be used to automatically "
        "reject or select candidates."
    )
